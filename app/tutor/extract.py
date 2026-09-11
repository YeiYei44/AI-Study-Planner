"""Extract text from uploaded material, one (page_ref, text) section per
page/slide/paragraph-group so citations stay accurate down to where the
text actually came from.
"""

from __future__ import annotations

from pathlib import Path


class UnsupportedFileType(ValueError):
    pass


def extract_sections(path: Path) -> list[tuple[str, str]]:
    """Returns [(page_ref, text), ...] — one entry per page/slide/section.
    An empty section (no extractable text) is dropped."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        sections = _extract_pdf(path)
    elif suffix == ".pptx":
        sections = _extract_pptx(path)
    elif suffix == ".docx":
        sections = _extract_docx(path)
    elif suffix in (".txt", ".md"):
        sections = _extract_text(path)
    else:
        raise UnsupportedFileType(
            f"No extractor for {suffix!r}. Supported: .pdf, .pptx, .docx, .txt, .md"
        )
    return [(ref, text) for ref, text in sections if text.strip()]


def _extract_pdf(path: Path) -> list[tuple[str, str]]:
    import pypdf

    reader = pypdf.PdfReader(str(path))
    return [
        (f"page {i + 1}", page.extract_text() or "")
        for i, page in enumerate(reader.pages)
    ]


def _extract_pptx(path: Path) -> list[tuple[str, str]]:
    import pptx

    prs = pptx.Presentation(str(path))
    sections = []
    for i, slide in enumerate(prs.slides):
        texts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame
        ]
        sections.append((f"slide {i + 1}", "\n".join(t for t in texts if t)))
    return sections


def _extract_docx(path: Path) -> list[tuple[str, str]]:
    import docx

    doc = docx.Document(str(path))
    # docx has no stored page concept (pagination is a rendering detail,
    # not part of the file) — group paragraphs into fixed-size "sections"
    # instead of faking page numbers.
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    group_size = 15
    sections = []
    for i in range(0, len(paras), group_size):
        section_num = i // group_size + 1
        sections.append((f"section {section_num}", "\n".join(paras[i : i + group_size])))
    return sections


def _extract_text(path: Path) -> list[tuple[str, str]]:
    text = path.read_text(errors="replace")
    return [("", text)]
